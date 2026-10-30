"""PoC 리포트 점검 (API 호출·정답 없음).  python -m pytest -q"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from bridgecore.constants import AXES, HOLD, NA, PRESERVED, VIOLATED  # noqa: E402
from bridgecore.io_utils import write_json, write_jsonl  # noqa: E402
from bridgecore.poc_report import (build_rows, highlight, make_poc_report, summarize,  # noqa: E402
                                   tally_review)


def _out(iid, judg: dict, sev: dict | None = None, errors=None, evaluable=True):
    sev = sev or {}
    return {"item_id": iid, "evaluable": evaluable, "evaluable_note": None if evaluable else "원문이 없음",
            "errors": errors or [],
            "axes": {a: {"judgment": judg.get(a, PRESERVED), "severity": sev.get(a),
                         "error_ids": [e["id"] for e in (errors or [])] if judg.get(a) == VIOLATED else [],
                         "reason": f"{a} 근거"} for a in AXES}}


def _fixture(tmp: Path):
    items = [
        {"item_id": "X1", "src": "The regents met.", "tgt": "섭정들이 모였다."},
        {"item_id": "X2", "src": "Close the door.", "tgt": "문을 닫아라."},
        {"item_id": "X3", "src": "Hi.", "tgt": "안녕."},
        {"item_id": "X4", "src": "Hello <b>", "tgt": "안녕 <b>"},
        {"item_id": "X5", "src": "No output.", "tgt": "출력 없음."},
    ]
    err = [{"id": "E001", "src_span": "regents", "tgt_span": "섭정들", "description": "기관이 바뀜"}]
    recs = [
        {"item_id": "X1", "model": "claude-haiku-5-5", "usage": None,
         "output": _out("X1", {"Semantic": VIOLATED, "Cultural": VIOLATED, "Deontic": NA},
                        {"Semantic": "major", "Cultural": "critical"}, err)},
        {"item_id": "X2", "model": "claude-haiku-5-5", "usage": None,
         "output": _out("X2", {"Register": HOLD, "Deontic": NA})},
        {"item_id": "X3", "model": "claude-haiku-5-5", "usage": None,
         "output": _out("X3", {"Deontic": NA, "Cultural": NA})},
        {"item_id": "X4", "model": "claude-haiku-5-5", "usage": None,
         "output": _out("X4", {}, evaluable=False)},
        {"item_id": "X5", "model": "claude-haiku-5-5", "usage": None, "output": None, "errors": ["타임아웃"]},
    ]
    run = tmp / "runs" / "demo"
    run.mkdir(parents=True)
    write_jsonl(run / "outputs.jsonl", recs)
    write_json(run / "meta.json", {"model": "claude-haiku-5-5", "prompt_version": "v0",
                                   "items": [i["item_id"] for i in items]})
    return items, recs, run


def test_rows_and_summary(tmp_path):
    items, recs, _ = _fixture(tmp_path)
    rows = build_rows({r["item_id"]: r for r in recs}, items)
    assert [r["status"] for r in rows] == ["ok", "ok", "ok", "not_evaluable", "no_output"]
    assert rows[0]["flagged"] == ["Semantic", "Cultural"] and rows[0]["max_sev"] == "critical"
    assert rows[1]["held"] == ["Register"] and not rows[1]["flagged"]
    s = summarize(rows)
    assert (s["n"], s["n_ok"], s["n_flagged"], s["n_held_only"], s["n_clean"]) == (5, 3, 1, 1, 1)
    assert (s["n_not_evaluable"], s["n_no_output"]) == (1, 1)
    assert s["per_axis"]["Cultural"]["critical"] == 1 and s["per_axis"]["Deontic"][NA] == 3


def test_highlight_overlap_and_escape():
    h, found = highlight("a <b> washed a buggy", [("E001", "washed a buggy"), ("E002", "a buggy"), ("E003", "없음")])
    assert found == {"E001", "E002"}
    assert "&lt;b&gt;" in h
    # 두 구간이 같은 자리에서 끝나므로 번호는 끝에 한 번만 붙는다
    assert h.count("<sup>") == 1 and "<sup>E001 E002</sup>" in h
    h2, _ = highlight("The Regents met", [("E001", "regents")])  # 대소문자 무시
    assert "<mark" in h2


def test_make_report_and_tally_roundtrip(tmp_path):
    from openpyxl import load_workbook
    items, _, run = _fixture(tmp_path)
    html_path, xlsx_path = make_poc_report(run, items, tmp_path / "reports")
    page = html_path.read_text(encoding="utf-8")
    assert "섭정들" in page and "X4" in page and "엔진 출력 없음" in page
    assert "Hello &lt;b&gt;" in page  # 원문 HTML 이스케이프

    wb = load_workbook(xlsx_path)
    flag = wb["엔진 표시"]
    rows = list(flag.iter_rows(min_row=2, values_only=True))
    # X1의 훼손 2칸 + X2의 보류 1칸, 심각한 문장이 먼저
    assert [(r[0], r[3], r[4]) for r in rows] == [("X1", "의미", VIOLATED), ("X1", "문화", VIOLATED),
                                                  ("X2", "관계와 격식", HOLD)]
    flag["I2"], flag["I3"], flag["I4"] = "맞음", "틀림", "맞음"
    full = wb["전체 문장"]
    assert full.max_row == 6
    full["E4"] = "의미, 문화"  # X3에서 놓친 오류
    guide = wb["안내"]
    hdr = next(c.row for c in guide["A"] if c.value == "검토자")
    guide.cell(hdr + 1, 1, "R001"); guide.cell(hdr + 1, 2, "10:00"); guide.cell(hdr + 1, 3, "10:40")
    wb.save(xlsx_path)

    t = tally_review(xlsx_path)
    assert t["flags_total"]["맞음"] == 1 and t["flags_total"]["틀림"] == 1
    assert t["flags_total"]["precision_lenient"] == 0.5
    assert t["holds"] == {"맞음": 1}
    assert t["n_items"] == 5 and t["n_items_with_missed"] == 1
    assert t["missed_by_axis"] == {"Semantic": 1, "Cultural": 1}
    assert t["time_log"][0]["검토자"] == "R001"


# ---------- 가이드라인 효과 비교 ----------
def test_generic_validate_and_prf():
    from bridgecore.baseline import validate
    from bridgecore.guideline_effect import _prf, bootstrap_diff
    assert validate({"errors": []}) == []
    assert validate({"errors": [{"src_span": "a", "tgt_span": "b", "description": "의미가 바뀜", "severity": "big"}]})
    gold = {"A": True, "B": False, "C": True, "D": False}
    m = _prf({"A": True, "B": True, "C": False, "D": False}, gold, list(gold))
    assert (m["tp"], m["fp"], m["fn"]) == (1, 1, 1) and abs(m["f1"] - 0.5) < 1e-9
    obs, lo, hi, _ = bootstrap_diff([dict(gold)], [{"A": True, "B": True, "C": False, "D": False}], gold, list(gold), n=200)
    assert abs(obs - 0.5) < 1e-9 and lo <= obs <= hi
