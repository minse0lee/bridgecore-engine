"""PoC 리포트: 엔진 판정을 사람이 읽는 보고서와 검토용 엑셀로 만든다.

정답(gold)을 쓰지 않는다. 본검수·책 번역처럼 정답이 없는 문장에 엔진을 돌린 뒤,
  1) <이름>_poc.html        — 요약, 축별 건수, 문장별 오류(구간 강조·심각도·근거)
  2) <이름>_poc_review.xlsx — 검토자가 엔진 판정이 맞는지 표시하고, 놓친 오류와 검토 시간을 적는 시트
를 만든다. 검토가 끝난 엑셀은 tally_review()로 집계한다.
"""
from __future__ import annotations

import html
from collections import Counter
from datetime import datetime
from pathlib import Path

from .constants import AXES, HOLD, NA, SEVERITIES, SEVERITY_RANK, VIOLATED
from .io_utils import read_json, write_json

AXIS_KO = {
    "Semantic": "의미", "Textual": "텍스트 구조", "Context": "발화 기능",
    "Deontic": "당위성", "Register": "관계와 격식", "Cultural": "문화",
}
AXIS_DESC = {
    "Semantic": "누가 무엇을 어떻게 했는지, 수량·시간·부정·조건·확실성이 유지되는가",
    "Textual": "지시 대상, 절 사이의 연결 관계, 초점 표현이 유지되는가",
    "Context": "진술·지시·질문·경고·제안 같은 발화 기능이 유지되는가",
    "Deontic": "의무·허가·금지와 그 대상·조건·예외가 유지되는가 (해당 문장만)",
    "Register": "관계에 맞는 높임·호칭·격식과 화자의 태도가 유지되는가",
    "Cultural": "제도·관습·단위·표기 같은 문화 고유 대상이 제대로 옮겨졌는가 (해당 문장만)",
}
SEV_KO = {"minor": "경미", "major": "중대", "critical": "치명"}
REVIEW_CHOICES = ["맞음", "틀림", "애매"]


# ───────────────────────── 집계 ─────────────────────────

def _max_sev(axes: dict) -> str | None:
    sevs = [a.get("severity") for a in axes.values() if a.get("judgment") == VIOLATED and a.get("severity")]
    return max(sevs, key=SEVERITY_RANK.get) if sevs else None


def build_rows(records: dict, items: list[dict]) -> list[dict]:
    """문장 하나 = 행 하나. items 순서를 따르되, 출력이 없는 문항은 status='no_output'."""
    rows = []
    for it in items:
        rec = records.get(it["item_id"])
        out = (rec or {}).get("output")
        row = {"item": it, "rec": rec, "out": out}
        if not out:
            row.update(status="no_output", flagged=[], held=[], max_sev=None)
        elif not out.get("evaluable", True):
            row.update(status="not_evaluable", flagged=[], held=[], max_sev=None)
        else:
            axes = out["axes"]
            row.update(status="ok",
                       flagged=[a for a in AXES if axes.get(a, {}).get("judgment") == VIOLATED],
                       held=[a for a in AXES if axes.get(a, {}).get("judgment") == HOLD],
                       max_sev=_max_sev(axes))
        rows.append(row)
    return rows


def summarize(rows: list[dict]) -> dict:
    ok = [r for r in rows if r["status"] == "ok"]
    per_axis = {a: Counter() for a in AXES}
    for r in ok:
        for a in AXES:
            ax = r["out"]["axes"].get(a, {})
            j = ax.get("judgment")
            per_axis[a][j] += 1
            if j == VIOLATED:
                per_axis[a][ax.get("severity") or "심각도 없음"] += 1
    return {
        "n": len(rows), "n_ok": len(ok),
        "n_flagged": sum(1 for r in ok if r["flagged"]),
        "n_held_only": sum(1 for r in ok if r["held"] and not r["flagged"]),
        "n_clean": sum(1 for r in ok if not r["flagged"] and not r["held"]),
        "n_not_evaluable": sum(1 for r in rows if r["status"] == "not_evaluable"),
        "n_no_output": sum(1 for r in rows if r["status"] == "no_output"),
        "by_max_sev": Counter(r["max_sev"] for r in ok if r["flagged"]),
        "per_axis": per_axis,
    }


def _order_key(r: dict):
    # 심각한 문장 먼저, 다음 보류만 있는 문장, 그다음 표시 없는 문장. 같은 묶음 안에서는 원래 순서.
    if r["flagged"]:
        return (0, -SEVERITY_RANK.get(r["max_sev"], -1))
    if r["held"] or r["status"] != "ok":
        return (1, 0)
    return (2, 0)


# ───────────────────────── 구간 강조 ─────────────────────────

def _find(text: str, span: str) -> int:
    if not span:
        return -1
    i = text.find(span)
    return i if i >= 0 else text.lower().find(span.lower())


def highlight(text: str, spans: list[tuple[str, str]]) -> tuple[str, set[str]]:
    """spans: [(오류ID, 구간)] → (강조된 HTML, 문장에서 찾은 오류ID 집합). 겹치는 구간도 처리한다."""
    marks: list[set[str]] = [set() for _ in text]
    ends: dict[int, list[str]] = {}
    found = set()
    for eid, span in spans:
        i = _find(text, span)
        if i < 0:
            continue
        found.add(eid)
        for k in range(i, i + len(span)):
            marks[k].add(eid)
        ends.setdefault(i + len(span), []).append(eid)
    out, k = [], 0
    while k < len(text):
        ids = marks[k]
        j = k
        while j < len(text) and marks[j] == ids:
            j += 1
        chunk = html.escape(text[k:j])
        if ids:
            # 오류 번호는 구간이 끝나는 자리에만 붙인다 (겹친 구간이 잘게 나뉘어도 번호가 반복되지 않게)
            sup = " ".join(sorted(x for x in ends.get(j, []) if x in ids))
            out.append(f'<mark title="{" ".join(sorted(ids))}">{chunk}'
                       + (f"<sup>{html.escape(sup)}</sup>" if sup else "") + "</mark>")
        else:
            out.append(chunk)
        k = j
    return "".join(out), found


# ───────────────────────── HTML ─────────────────────────

CSS = """
:root{--bg:#fbfaf7;--card:#fff;--ink:#1f2328;--mute:#5f6670;--line:#e4e1da;--accent:#2f5d8a;
--crit:#a3242a;--crit-bg:#fbe9ea;--maj:#b45309;--maj-bg:#fdf0e1;--min:#5b6b2a;--min-bg:#eef2e0;
--hold:#5b4b8a;--hold-bg:#efecf7;--mark:#ffe8a3}
@media (prefers-color-scheme:dark){:root{--bg:#16181b;--card:#1e2125;--ink:#e6e6e3;--mute:#a0a6ad;--line:#33373c;
--accent:#8db4dc;--crit:#f2a0a4;--crit-bg:#3a1f21;--maj:#f3b770;--maj-bg:#3a2c1b;--min:#c4d28f;--min-bg:#2a2f1d;
--hold:#c2b6ea;--hold-bg:#2a2638;--mark:#5c4a14}}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.65 "Pretendard","Apple SD Gothic Neo","Malgun Gothic","Noto Sans KR",system-ui,sans-serif}
main{max-width:980px;margin:0 auto;padding:32px 16px 80px}
h1{font-size:24px;margin:0 0 4px} h2{font-size:18px;margin:36px 0 12px;padding-bottom:6px;border-bottom:1px solid var(--line)}
.meta{color:var(--mute);font-size:13px}
.notice p{margin:4px 0}
.notice{background:var(--card);border:1px solid var(--line);border-left:4px solid var(--accent);padding:12px 14px;margin:18px 0;font-size:14px}
.tiles{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px}
.tile{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:12px 14px}
.tile b{display:block;font-size:24px;font-variant-numeric:tabular-nums}
.tile span{color:var(--mute);font-size:13px}
table{border-collapse:collapse;width:100%;background:var(--card);font-size:14px}
th,td{border:1px solid var(--line);padding:6px 8px;text-align:left;vertical-align:top}
th{background:var(--bg);font-weight:600} td.n{text-align:right;font-variant-numeric:tabular-nums}
.tablewrap{overflow-x:auto}
.filters{display:flex;flex-wrap:wrap;gap:6px;align-items:center;margin:8px 0 16px;font-size:14px}
.filters button{border:1px solid var(--line);background:var(--card);color:var(--ink);border-radius:999px;padding:4px 12px;cursor:pointer;font:inherit}
.filters button.on{background:var(--accent);border-color:var(--accent);color:#fff}
.filters label{margin-left:8px;color:var(--mute)}
.card{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:14px 16px;margin:0 0 12px}
.card header{display:flex;flex-wrap:wrap;gap:6px 10px;align-items:baseline;margin-bottom:8px}
.card header .id{font-weight:700} .card header .doc{color:var(--mute);font-size:13px}
.pair{display:grid;grid-template-columns:52px 1fr;gap:4px 10px;margin:6px 0 10px}
.pair .k{color:var(--mute);font-size:13px;padding-top:2px}
mark{background:var(--mark);color:inherit;padding:0 1px;border-radius:2px}
mark sup{font-size:10px;color:var(--mute);margin-left:1px}
.badge{display:inline-block;font-size:12px;font-weight:600;border-radius:4px;padding:1px 7px;white-space:nowrap}
.critical{color:var(--crit);background:var(--crit-bg)} .major{color:var(--maj);background:var(--maj-bg)}
.minor{color:var(--min);background:var(--min-bg)} .hold{color:var(--hold);background:var(--hold-bg)}
.plain{color:var(--mute);background:var(--bg);border:1px solid var(--line)}
.errs{margin:6px 0 8px;padding-left:0;list-style:none;font-size:14px}
.errs li{margin:4px 0} .errs code{font-family:inherit;font-weight:600}
.axrow{display:grid;grid-template-columns:120px 1fr;gap:4px 10px;padding:6px 0;border-top:1px dashed var(--line);font-size:14px}
.axrow .ax{font-weight:600}
details{margin-top:8px;font-size:13px;color:var(--mute)} details p{white-space:pre-wrap;margin:6px 0}
.hidden{display:none}
@media (max-width:560px){.pair{grid-template-columns:1fr}.axrow{grid-template-columns:1fr}}
"""

JS = """
(function(){
  var axis='all', sev={critical:1,major:1,minor:1,none:1}, showClean=false;
  var cards=[].slice.call(document.querySelectorAll('.card'));
  function apply(){
    var shown=0;
    cards.forEach(function(c){
      var fl=(c.dataset.axsev||'').split(' ').filter(Boolean).map(function(x){return x.split(':')});
      var held=(c.dataset.held||'').split(' ').filter(Boolean);
      var ok;
      if(c.dataset.status!=='ok'){ ok=axis==='all'; }
      else if(!fl.length&&!held.length){ ok=showClean&&axis==='all'; }
      else {
        ok=fl.some(function(p){return (axis==='all'||p[0]===axis)&&sev[p[1]]})
          ||held.some(function(a){return axis==='all'||a===axis});
      }
      c.classList.toggle('hidden',!ok); if(ok) shown++;
    });
    document.getElementById('shown').textContent=shown;
  }
  document.querySelectorAll('[data-axis]').forEach(function(b){
    b.addEventListener('click',function(){
      axis=b.dataset.axis;
      document.querySelectorAll('[data-axis]').forEach(function(x){x.classList.toggle('on',x===b)});
      apply();
    });
  });
  document.querySelectorAll('[data-sevbox]').forEach(function(x){
    x.addEventListener('change',function(){ sev[x.dataset.sevbox]=x.checked?1:0; apply(); });
  });
  document.getElementById('clean').addEventListener('change',function(e){ showClean=e.target.checked; apply(); });
  apply();
})();
"""


def _sev_badge(sev: str | None) -> str:
    if not sev:
        return '<span class="badge plain">심각도 없음</span>'
    return f'<span class="badge {sev}">{SEV_KO.get(sev, sev)}</span>'


def _card(r: dict) -> str:
    it, out = r["item"], r["out"]
    e = html.escape
    doc = " · ".join(x for x in (it.get("genre"), it.get("title")) if x)
    head = [f'<span class="id">{e(it["item_id"])}</span>']
    for a in r["flagged"]:
        sev = out["axes"][a].get("severity")
        head.append(f'<span class="badge {sev or "plain"}">{AXIS_KO[a]} · {SEV_KO.get(sev, "심각도 없음")}</span>')
    for a in r["held"]:
        head.append(f'<span class="badge hold">{AXIS_KO[a]} · 판단보류</span>')
    if r["status"] == "not_evaluable":
        head.append('<span class="badge plain">평가 불가</span>')
    if r["status"] == "no_output":
        head.append('<span class="badge plain">엔진 출력 없음</span>')
    if doc:
        head.append(f'<span class="doc">{e(doc)}</span>')

    tgt = (it.get("tgt") or "").strip()
    errors = (out or {}).get("errors", [])
    src_html, src_found = highlight(it.get("src", ""), [(x["id"], x.get("src_span", "")) for x in errors])
    tgt_html, tgt_found = highlight(tgt, [(x["id"], x.get("tgt_span", "")) for x in errors])
    body = [f'<div class="pair"><div class="k">원문</div><div>{src_html}</div>'
            f'<div class="k">번역</div><div>{tgt_html or "<i>(문단 번역에서 엔진이 대응 구간을 찾음)</i>"}</div></div>']

    if r["status"] == "not_evaluable":
        body.append(f'<p>엔진이 이 문장은 평가할 수 없다고 했습니다: {e(out.get("evaluable_note") or "")}</p>')
    if r["status"] == "no_output":
        errs = "; ".join((r["rec"] or {}).get("errors", []) or []) if r["rec"] else "실행 기록 없음"
        body.append(f'<p>엔진 출력이 없습니다. {e(errs)}</p>')

    if errors:
        used = {}
        for a in AXES:
            for eid in (out["axes"].get(a, {}).get("error_ids") or []):
                used.setdefault(eid, []).append(AXIS_KO[a])
        lis = []
        for x in errors:
            where = ""
            if x["id"] not in src_found and x["id"] not in tgt_found:
                where = " <span class=\"meta\">(문장 앞뒤 문맥에 있는 구간)</span>"
            axes_txt = f' <span class="meta">→ {", ".join(used[x["id"]])}</span>' if x["id"] in used else ""
            lis.append(f'<li><code>{e(x["id"])}</code> “{e(x.get("src_span") or "")}” → “{e(x.get("tgt_span") or "")}”'
                       f'{where}{axes_txt}<br>{e(x.get("description") or "")}</li>')
        body.append(f'<ul class="errs">{"".join(lis)}</ul>')

    for a in r["flagged"] + r["held"]:
        ax = out["axes"][a]
        badge = _sev_badge(ax.get("severity")) if ax["judgment"] == VIOLATED else '<span class="badge hold">판단보류</span>'
        body.append(f'<div class="axrow"><div class="ax">{AXIS_KO[a]} {badge}</div><div>{e(ax.get("reason") or "")}</div></div>')

    if out and r["status"] == "ok" and not r["flagged"] and not r["held"]:
        body.append('<p class="meta">여섯 축 모두 보존 또는 해당 없음(N/A)으로 판정했습니다.</p>')

    if it.get("src_paragraph") or it.get("tgt_paragraph"):
        body.append(f'<details><summary>문단 전체 보기</summary><p>{e(it.get("src_paragraph") or "")}</p>'
                    f'<p>{e(it.get("tgt_paragraph") or "")}</p></details>')

    axsev = " ".join(f'{a}:{out["axes"][a].get("severity") or "none"}' for a in r["flagged"])
    attrs = f'data-status="{r["status"]}" data-axsev="{axsev}" data-held="{" ".join(r["held"])}"'
    return f'<article class="card" {attrs}><header>{"".join(head)}</header>{"".join(body)}</article>'


def render_html(rows: list[dict], summary: dict, info: dict) -> str:
    e = html.escape
    s = summary
    pct = lambda k: f'{s[k] / s["n_ok"] * 100:.0f}%' if s["n_ok"] else "-"
    tiles = [
        (s["n"], "전체 문장"),
        (s["n_flagged"], f"훼손 표시 문장 ({pct('n_flagged')})"),
        (s["by_max_sev"].get("critical", 0) + s["by_max_sev"].get("major", 0), "그중 중대·치명 포함"),
        (s["n_held_only"], "판단보류만 있는 문장"),
        (s["n_clean"], f"표시 없는 문장 ({pct('n_clean')})"),
    ]
    if s["n_not_evaluable"] or s["n_no_output"]:
        tiles.append((s["n_not_evaluable"] + s["n_no_output"], "평가 불가·출력 없음"))
    tiles_html = "".join(f'<div class="tile"><b>{n}</b><span>{e(t)}</span></div>' for n, t in tiles)

    ax_rows = []
    for a in AXES:
        c = s["per_axis"][a]
        ax_rows.append(
            f'<tr><td><b>{AXIS_KO[a]}</b> <span class="meta">{a}</span><br><span class="meta">{AXIS_DESC[a]}</span></td>'
            f'<td class="n">{c[VIOLATED]}</td>' + "".join(f'<td class="n">{c[x]}</td>' for x in reversed(SEVERITIES))
            + f'<td class="n">{c[HOLD]}</td><td class="n">{c[NA]}</td></tr>')
    axis_table = ('<div class="tablewrap"><table><thead><tr><th>축</th><th>훼손</th><th>치명</th><th>중대</th><th>경미</th>'
                  '<th>판단보류</th><th>해당 없음</th></tr></thead><tbody>' + "".join(ax_rows) + '</tbody></table></div>')

    cost = info.get("cost")
    cost_html = ""
    if cost and cost.get("usd") is not None:
        per = cost["usd_per_item"] or 0
        mode = "Batch" if info.get("via_batch") else "실시간 호출"
        cost_html = (f'<p class="meta">API 비용: 이번 실행 ${cost["usd"]:.2f} ({mode}), 문장당 ${per:.4f}, '
                     f'같은 조건으로 1만 문장이면 약 ${per * 10000:,.0f}')
        bc = info.get("cost_batch")
        if not info.get("via_batch") and bc and bc.get("usd_per_item") is not None:
            cost_html += f', Batch로 돌리면 약 ${bc["usd_per_item"] * 10000:,.0f}'
        cost_html += '.</p>'

    filters = ['<button class="on" data-axis="all">전체</button>'] + [
        f'<button data-axis="{a}">{AXIS_KO[a]}</button>' for a in AXES]
    sevboxes = "".join(f'<label><input type="checkbox" data-sevbox="{x}" checked> {SEV_KO[x]}</label>'
                       for x in reversed(SEVERITIES))
    cards = "".join(_card(r) for r in sorted(rows, key=_order_key))
    note = f'<p>{e(info["note"])}</p>' if info.get("note") else ""

    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{e(info["title"])}</title><style>{CSS}</style></head><body><main>
<h1>{e(info["title"])}</h1>
<div class="meta">실행 {e(info["run"])} · 모델 {e(info["model"])} · 프롬프트 {e(info["prompt_version"])} · 생성 {e(info["created"])}</div>
<div class="notice">{note}<p>이 보고서는 사람이 검토하기 전의 엔진 판정입니다. 훼손 표시는 검토자가 먼저 확인할 후보이고 확정 판정이 아닙니다.
엔진은 번역문을 여섯 축으로 나누어 보고, 축마다 보존·훼손·해당 없음(N/A)·판단보류 가운데 하나를 고릅니다.
훼손이면 경미·중대·치명 가운데 심각도를 붙이고, 원문과 번역문에서 문제가 된 구간을 짚습니다.</p></div>
<h2>요약</h2><div class="tiles">{tiles_html}</div>
<h2>축별 판정 수</h2>{axis_table}{cost_html}
<h2>문장별 결과</h2>
<div class="filters">{"".join(filters)}<span style="flex-basis:100%;height:0"></span>심각도:{sevboxes}
<label><input type="checkbox" id="clean"> 표시 없는 문장도 보기</label>
<span class="meta" style="margin-left:auto"><span id="shown"></span>문장 표시 중</span></div>
{cards}
</main><script>{JS}</script></body></html>"""


# ───────────────────────── 검토용 엑셀 ─────────────────────────

def write_review_xlsx(rows: list[dict], path: Path, info: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="E8EEF5")
    input_fill = PatternFill("solid", fgColor="FFF7DB")
    wrap = Alignment(wrap_text=True, vertical="top")
    bold = Font(bold=True)

    # 안내
    ws = wb.active
    ws.title = "안내"
    guide = [
        info["title"],
        f'실행 {info["run"]} · 모델 {info["model"]} · 프롬프트 {info["prompt_version"]}',
        "",
        "노란 칸만 입력해 주세요.",
        "1. '엔진 표시' 시트: 엔진이 훼손 또는 판단보류로 표시한 칸입니다. 표시가 맞으면 '맞음', 틀리면 '틀림', "
        "판단이 어려우면 '애매'를 고르고, 틀림·애매일 때 검토자 판정과 이유를 적습니다.",
        "2. '전체 문장' 시트: 엔진이 놓친 오류가 있으면 그 축을 쉼표로 적고(예: 의미, 문화) 짧게 설명합니다. 없으면 비워 둡니다.",
        "3. 아래 시간 칸에 검토를 시작·끝낸 시각과 쉰 시간(분)을 적습니다. 엔진 결과를 참고했을 때 시간이 얼마나 드는지 보려는 것입니다.",
        "",
        "축 설명",
    ] + [f"{AXIS_KO[a]} ({a}): {AXIS_DESC[a]}" for a in AXES]
    for i, line in enumerate(guide, 1):
        ws.cell(i, 1, line).alignment = Alignment(wrap_text=True)
    ws["A1"].font = Font(bold=True, size=14)
    t0 = len(guide) + 2
    for j, h in enumerate(["검토자", "시작 시각", "끝난 시각", "쉰 시간(분)", "검토한 문장 수"], 1):
        c = ws.cell(t0, j, h)
        c.font, c.fill = bold, head_fill
        for k in range(1, 4):
            ws.cell(t0 + k, j).fill = input_fill
    ws.column_dimensions["A"].width = 110
    for col in "BCDE":
        ws.column_dimensions[col].width = 16

    # 엔진 표시
    ws = wb.create_sheet("엔진 표시")
    cols = ["문항ID", "원문", "번역문", "축", "엔진 판정", "엔진 심각도", "오류 구간", "엔진 근거",
            "엔진 표시가 맞나", "검토자 판정", "검토자 심각도", "검토 의견"]
    widths = [9, 45, 45, 12, 10, 10, 40, 50, 12, 11, 11, 40]
    ws.append(cols)
    for row in sorted(rows, key=_order_key):
        out = row["out"]
        if not out:
            continue
        err_by_id = {x["id"]: x for x in out.get("errors", [])}
        for a in row["flagged"] + row["held"]:
            ax = out["axes"][a]
            spans = "\n".join(f'{eid}: “{err_by_id[eid].get("src_span", "")}” → “{err_by_id[eid].get("tgt_span", "")}”'
                              for eid in (ax.get("error_ids") or []) if eid in err_by_id)
            ws.append([row["item"]["item_id"], row["item"].get("src", ""), row["item"].get("tgt", ""),
                       AXIS_KO[a], ax["judgment"], SEV_KO.get(ax.get("severity"), ""), spans,
                       ax.get("reason", ""), None, None, None, None])
    _finish_sheet(ws, widths, input_cols=[9, 10, 11, 12], head_fill=head_fill, input_fill=input_fill, wrap=wrap)
    n = ws.max_row
    if n >= 2:
        _validation(ws, DataValidation, REVIEW_CHOICES, f"I2:I{n}")
        _validation(ws, DataValidation, ["보존", "훼손", "N/A", "판단보류"], f"J2:J{n}")
        _validation(ws, DataValidation, [SEV_KO[x] for x in SEVERITIES], f"K2:K{n}")

    # 전체 문장
    ws = wb.create_sheet("전체 문장")
    ws.append(["문항ID", "원문", "번역문", "엔진 표시", "놓친 오류 축", "놓친 오류 설명"])
    for row in rows:
        marks = [f"{AXIS_KO[a]}({SEV_KO.get(row['out']['axes'][a].get('severity'), '-')})" for a in row["flagged"]]
        marks += [f"{AXIS_KO[a]}(판단보류)" for a in row["held"]]
        if row["status"] == "not_evaluable":
            marks = ["평가 불가"]
        elif row["status"] == "no_output":
            marks = ["엔진 출력 없음"]
        ws.append([row["item"]["item_id"], row["item"].get("src", ""), row["item"].get("tgt", ""),
                   ", ".join(marks) or "표시 없음", None, None])
    _finish_sheet(ws, [9, 50, 50, 26, 18, 45], input_cols=[5, 6], head_fill=head_fill,
                  input_fill=input_fill, wrap=wrap)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _finish_sheet(ws, widths, input_cols, head_fill, input_fill, wrap):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
        c = ws.cell(1, j)
        c.font, c.fill, c.alignment = Font(bold=True), head_fill, wrap
    for r in range(2, ws.max_row + 1):
        for j in range(1, len(widths) + 1):
            c = ws.cell(r, j)
            c.alignment = wrap
            if j in input_cols:
                c.fill = input_fill
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def _validation(ws, DataValidation, choices, ref):
    dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(ref)


# ───────────────────────── 검토 결과 집계 ─────────────────────────

def tally_review(xlsx_path: Path) -> dict:
    """검토가 끝난 엑셀을 읽어 엔진 표시 정확도와 놓친 오류를 센다.
    precision_strict: 맞음 / (맞음+틀림+애매),  precision_lenient: 맞음 / (맞음+틀림)  — 훼손 표시만 대상."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, data_only=True)
    ko2en = {v: k for k, v in AXIS_KO.items()}

    flags = {a: Counter() for a in AXES}
    holds = Counter()
    for r in wb["엔진 표시"].iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        axis, judgment, verdict = ko2en.get(r[3], r[3]), r[4], (r[8] or "").strip() or "미입력"
        if judgment == VIOLATED:
            flags.setdefault(axis, Counter())[verdict] += 1
        else:
            holds[verdict] += 1

    missed = Counter()
    n_items = n_with_miss = 0
    for r in wb["전체 문장"].iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        n_items += 1
        axes = [x.strip() for x in str(r[4] or "").replace("，", ",").split(",") if x.strip()]
        if axes:
            n_with_miss += 1
        for x in axes:
            missed[ko2en.get(x, x)] += 1

    def prec(c: Counter):
        ok, bad, unsure = c["맞음"], c["틀림"], c["애매"]
        return {"맞음": ok, "틀림": bad, "애매": unsure, "미입력": c["미입력"],
                "precision_strict": ok / (ok + bad + unsure) if ok + bad + unsure else None,
                "precision_lenient": ok / (ok + bad) if ok + bad else None}

    total = sum(flags.values(), Counter())
    # 시간 칸은 적는 형식이 사람마다 달라 계산하지 않고 그대로 옮긴다
    times = []
    ws = wb["안내"]
    hdr = next((i for i, r in enumerate(ws.iter_rows(values_only=True), 1) if r and r[0] == "검토자"), None)
    if hdr:
        for r in ws.iter_rows(min_row=hdr + 1, max_row=hdr + 3, values_only=True):
            if any(v not in (None, "") for v in r[:5]):
                times.append({"검토자": r[0], "시작": str(r[1] or ""), "끝": str(r[2] or ""),
                              "쉰 시간(분)": r[3], "문장 수": r[4]})

    return {"flags_total": prec(total), "flags_by_axis": {a: prec(c) for a, c in flags.items() if sum(c.values())},
            "holds": dict(holds), "n_items": n_items, "n_items_with_missed": n_with_miss,
            "missed_by_axis": dict(missed), "time_log": times}


# ───────────────────────── 진입점 ─────────────────────────

def make_poc_report(run_dir: Path, items: list[dict], reports_dir: Path, *, title: str | None = None,
                    note: str | None = None, name: str | None = None) -> tuple[Path, Path]:
    from .report import _calls_for_run, _cost
    from .runner import latest_records

    run_dir = Path(run_dir)
    meta = read_json(run_dir / "meta.json")
    recs = latest_records(run_dir)
    want = set(meta.get("items") or recs)
    sel = [it for it in items if it["item_id"] in want]
    if not sel:
        raise SystemExit(f"{run_dir}: 문항 파일과 실행 결과의 item_id가 겹치지 않습니다.")
    rows = build_rows(recs, sel)
    summary = summarize(rows)
    ids = {it["item_id"] for it in sel}
    info = {
        "title": title or f"번역 검증 엔진 결과 — {run_dir.name}",
        "note": note, "run": run_dir.name, "model": meta.get("model", "?"),
        "prompt_version": meta.get("prompt_version", "?"),
        "created": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "cost": _cost(calls := _calls_for_run(run_dir, meta, recs, ids), len(ids)),
        "cost_batch": _cost([{**c, "via_batch": True} for c in calls], len(ids)),
        "via_batch": any((recs.get(i) or {}).get("via_batch") for i in ids),
    }
    base = name or run_dir.name
    reports_dir.mkdir(parents=True, exist_ok=True)
    html_path = reports_dir / f"{base}_poc.html"
    html_path.write_text(render_html(rows, summary, info), encoding="utf-8")
    xlsx_path = write_review_xlsx(rows, reports_dir / f"{base}_poc_review.xlsx", info)
    write_json(reports_dir / f"{base}_poc_summary.json", {
        **{k: v for k, v in summary.items() if k not in ("per_axis", "by_max_sev")},
        "by_max_sev": dict(summary["by_max_sev"]),
        "per_axis": {a: dict(c) for a, c in summary["per_axis"].items()},
        "cost_usd": (info["cost"] or {}).get("usd"), "run": info["run"], "model": info["model"],
        "prompt_version": info["prompt_version"],
    })
    print(f"문장 {summary['n']}개 · 훼손 표시 {summary['n_flagged']} · 판단보류만 {summary['n_held_only']} · "
          f"표시 없음 {summary['n_clean']}\n→ {html_path}\n→ {xlsx_path}")
    return html_path, xlsx_path
